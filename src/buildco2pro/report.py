"""Excel report: the pages of the dashboard as sheets, every number with its factor and source."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import __version__

HEAD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="30343B")


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = HEAD; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths or [18] * len(header)):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    return ws


def build_workbook(p, res, lib, checks=None) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Summary"
    ws.append([f"BuildCO2 Pro {__version__}: upfront carbon (EN 15978 A1-A5)"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Project", p.name]); ws.append(["Gross floor area (m2)", p.gfa_m2]); ws.append(["Factor set", lib.version])
    ws.append(["Sheet products", p.sheet_products]); ws.append([])
    ws.append(["Module", "tCO2e", "kgCO2e/m2 GFA"])
    for m, v in res.modules_t.items():
        ws.append([m, round(v, 1), round(v * 1e3 / p.gfa_m2, 1)])
    ws.append(["Total A1-A5", round(res.total_t, 1), round(res.per_m2_kg, 1)])
    for w in res.warnings:
        ws.append(["Warning", w])
    ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 60
    _sheet(wb, "A1-A3 by material", ["Material class", "Group", "Name", "Quantity", "Unit", "Factor", "Factor unit",
                                     "Mass (t)", "tCO2e", "Rows", "Source", "Tier"],
           [(l.material_class, l.group, l.name, round(l.quantity, 2), l.unit, l.factor, l.factor_unit, round(l.mass_t, 2),
             round(l.tCO2e, 2), l.rows, l.source, l.tier) for l in res.lines], [22, 18, 30, 12, 8, 10, 12, 10, 10, 8, 50, 6])
    cats = sorted({k.split("|")[0] for k in res.a1a3_by_category_group}); grps = list(res.a1a3_by_group)
    _sheet(wb, "A1-A3 by element", ["Category"] + grps + ["Total"],
           [[c] + [round(res.a1a3_by_category_group.get(f"{c}|{g}", 0.0), 1) for g in grps] + [round(res.a1a3_by_category.get(c, 0), 1)]
            for c in sorted(cats, key=lambda c: -res.a1a3_by_category.get(c, 0))], [24] + [14] * (len(grps) + 1))
    _sheet(wb, "A4 and A5", ["Module", "Item", "tCO2e"],
           [("A4", g, round(v, 2)) for g, v in res.a4_by_group_t.items()] + [("A5w", g, round(v, 2)) for g, v in res.a5w_by_group_t.items()]
           + [("A5a", c, round(v, 2)) for c, v in res.a5a_by_crane_t.items()] + [("A5r", "Spoil", round(res.modules_t["A5r"], 2))], [10, 28, 12])
    if res.storey_curve:
        _sheet(wb, "Storeys", ["Storeys", "A1-A3 (t)", "A4 (t)", "A5w (t)", "Total (t)"],
               [(r["storeys"], round(r["A1A3_t"], 1), round(r["A4_t"], 1), round(r["A5w_t"], 1), round(r["total_t"], 1)) for r in res.storey_curve])
    if res.scenarios:
        _sheet(wb, "Scenarios", ["Lever", "A1-A3 (t)", "Total (t)", "Saving (t)", "Saving (%)"],
               [(s["name"], round(s["A1A3_t"], 1), round(s["total_t"], 1), round(s["saving_t"], 1), round(s["saving_pct"], 1)) for s in res.scenarios])
    _sheet(wb, "Checks", ["Check", "Severity", "Finding", "Estimate", "Examples", "Action"],
           [(c["check"], c["severity"], c["finding"], c["estimate"], c["examples"], c["action"]) for c in (checks or [])], [7, 9, 60, 36, 40, 60])
    _sheet(wb, "Not classified", ["Category", "Material", "Reason", "Rows", "Volume (m3)", "Area (m2)"],
           [tuple(k.split("|")) + (u["rows"], round(u["volume_m3"], 2), round(u["area_m2"], 1)) for k, u in res.unclassified.items()], [24, 40, 14, 8, 12, 12])
    _sheet(wb, "Factors", ["Key", "Group", "Name", "Basis", "Value", "Unit", "Density", "Source", "Tier", "Notes"],
           [(f["key"], f["group"], f["name"], f["basis"], f["value"], f["unit"], f["density"], f["source"], f["tier"], f["notes"]) for f in lib.rows()],
           [22, 18, 30, 8, 8, 12, 8, 60, 5, 50])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
