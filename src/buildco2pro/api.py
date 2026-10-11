"""FastAPI service: JSON API and the browser dashboard (served from /).

Inventories are uploaded once (IFC, take-off CSV or Excel, Power BI file) and kept in memory under a token;
the dashboard then sends only the project inputs, so that changing a factor or a distance recalculates at once.
"""
from __future__ import annotations

import io
import tempfile
import uuid
from importlib import resources
from pathlib import Path

from fastapi import Body, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import __version__
from .carbon import assess
from .checks import run as run_checks
from .examples import example_inventory, example_project
from .factors import SETS, FactorLibrary
from .inventory import Inventory, from_ifc, from_pbix, from_takeoff_csv, from_takeoff_rows
from .models import ProjectInput
from .report import build_workbook

app = FastAPI(title="BuildCO2 Pro", version=__version__,
              description="Upfront embodied carbon (EN 15978 A1-A5) of buildings from IFC or a Revit take-off. Open source (MIT).")
WEB = Path(str(resources.files("buildco2pro") / "web"))
STORE: dict[str, dict] = {}          # token -> {inventory, files, name}
TMP = Path(tempfile.gettempdir()) / "buildco2pro"; TMP.mkdir(exist_ok=True)


class AssessRequest(BaseModel):
    project: ProjectInput = ProjectInput()
    token: str = "example"


def _inv(token) -> Inventory:
    if token == "example":
        STORE.setdefault("example", dict(inventory=example_inventory(), files=[], name="Example building"))
    if token not in STORE:
        raise HTTPException(404, "inventory not found: upload the model again")
    return STORE[token]["inventory"]


def _summary(inv: Inventory) -> dict:
    by_src, n_cls = {}, 0
    for r in inv.rows:
        by_src[r.source] = by_src.get(r.source, 0) + 1; n_cls += bool(r.material_class)
    storeys = sorted({r.storey for r in inv.rows if r.storey})
    return dict(rows=len(inv), classified=n_cls, sources=by_src, storeys=storeys)


@app.get("/api/health")
def health():
    return {"status": "ok", "version": __version__, "factor_sets": SETS}


@app.get("/api/defaults")
def defaults():
    return {"project": ProjectInput().model_dump(), "example": example_project().model_dump()}


@app.get("/api/factors")
def factors(set: str = "epic2019"):
    lib = FactorLibrary(set)
    return {"set": set, "version": lib.version, "materials": lib.rows(),
            "site": [s.__dict__ for s in lib.site.values()], "waste": lib.waste_rows}


@app.post("/api/upload")
async def upload(files: list[UploadFile] = File(...)):
    token = uuid.uuid4().hex[:12]; inv = Inventory(); kept = []
    for up in files:
        name = Path(up.filename or "model").name; data = await up.read()
        path = TMP / f"{token}_{name}"; path.write_bytes(data)
        low = name.lower()
        try:
            if low.endswith(".ifc"):
                inv.extend(from_ifc(path)); kept.append(str(path))
            elif low.endswith(".pbix"):
                inv.extend(from_pbix(path)[0])
            elif low.endswith(".csv"):
                head = data[:400].decode("utf-8", "ignore")
                inv.extend(Inventory.from_csv(data.decode("utf-8")) if "material_class" in head and "rule" in head
                           else from_takeoff_csv(path))
            elif low.endswith((".xlsx", ".xlsm")):
                inv.extend(_from_xlsx(path, name))
            else:
                raise HTTPException(400, f"{name}: use .ifc, .csv, .xlsx or .pbix")
        except HTTPException:
            raise
        except Exception as e:                                  # noqa: BLE001
            raise HTTPException(400, f"{name}: {e}")
    STORE[token] = dict(inventory=inv, files=kept, name=", ".join(f.filename or "" for f in files))
    return {"token": token, "summary": _summary(inv)}


def _from_xlsx(path, name):
    """Every sheet with the Revit take-off columns becomes rows; the sheet name is the element category."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True); inv = Inventory()
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        hdr = None
        for row in it:
            if row and any(isinstance(c, str) and c.strip().startswith("Material: Name") for c in row):
                hdr = [str(c).strip() if c else "" for c in row]; break
        if not hdr:
            continue
        recs = [dict(zip(hdr, r)) for r in it if r and any(r)]
        inv.extend(from_takeoff_rows(recs, source=f"{name}:{ws.title}", category=ws.title))
    return inv


@app.get("/api/inventory/{token}")
def inventory(token: str, q: str = "", limit: int = Query(300, le=5000), unclassified: bool = False):
    inv = _inv(token); ql = q.lower(); out = []
    for r in inv.rows:
        if unclassified and r.material_class:
            continue
        if ql and ql not in f"{r.category} {r.family_type} {r.material} {r.material_class} {r.storey}".lower():
            continue
        out.append(r.__dict__)
        if len(out) >= limit:
            break
    return {"summary": _summary(inv), "rows": out}


@app.get("/api/inventory/{token}.csv")
def inventory_csv(token: str):
    return Response(_inv(token).to_csv(), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=inventory.csv"})


@app.post("/api/assess")
def api_assess(req: AssessRequest = Body(...)):
    inv = _inv(req.token); lib = FactorLibrary(req.project.factor_set, req.project.factor_overrides)
    res = assess(req.project, inv, lib)
    return {"result": res.to_dict(), "checks": run_checks(inv, lib), "summary": _summary(inv), "factor_version": lib.version}


@app.post("/api/report.xlsx")
def api_report(req: AssessRequest = Body(...)):
    inv = _inv(req.token); lib = FactorLibrary(req.project.factor_set, req.project.factor_overrides)
    res = assess(req.project, inv, lib)
    data = build_workbook(req.project, res, lib, run_checks(inv, lib))
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=buildco2pro_report.xlsx"})


@app.post("/api/tagged.ifc")
def api_tagged(req: AssessRequest = Body(...)):
    """The first uploaded IFC file with a BuildCO2_Carbon property set on every assessed element."""
    from .ifcexport import tag_ifc
    st = STORE.get(req.token)
    if not st or not st["files"]:
        raise HTTPException(400, "no IFC file uploaded for this inventory")
    lib = FactorLibrary(req.project.factor_set, req.project.factor_overrides)
    out = TMP / f"{req.token}_tagged.ifc"
    tag_ifc(st["files"][0], out, st["inventory"], lib, req.project)
    return FileResponse(out, filename=Path(st["files"][0]).name.split("_", 1)[-1].replace(".ifc", "_carbon.ifc"))


@app.post("/api/figure/{name}")
def api_figure(name: str, req: AssessRequest = Body(...), fmt: str = "svg"):
    from . import figures as F
    if name not in F.FIGURES:
        raise HTTPException(404, f"figure {name!r}: choose from {list(F.FIGURES)}")
    inv = _inv(req.token); res = assess(req.project, inv)
    try:
        data = F.render(name, res, fmt, inv=inv)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return Response(data, media_type={"svg": "image/svg+xml", "pdf": "application/pdf", "png": "image/png"}[fmt])


@app.get("/", response_class=HTMLResponse)
def index():
    return (WEB / "index.html").read_text(encoding="utf-8")


app.mount("/static", StaticFiles(directory=str(WEB)), name="static")


def run(host="127.0.0.1", port=8000, open_browser=True):
    import threading
    import webbrowser
    import uvicorn
    if open_browser:
        threading.Timer(1.2, lambda: webbrowser.open(f"http://{host}:{port}")).start()
    uvicorn.run(app, host=host, port=port, log_level="warning")
